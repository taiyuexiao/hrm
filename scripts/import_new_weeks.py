#!/usr/bin/env python3
"""
安全导入新的周报数据（20260424, 20260518, 20260525, 20260529）到 SQLite。
核心原则：只操作目标周次，绝不碰存量数据。
"""
import json
import os
import re
import sqlite3
from datetime import datetime

import openpyxl

# ============== 配置 ==============
DB_PATH = 'backend/data/hr.db'
EXCEL_DIR = 'data/周报/部门周报'
TARGET_WEEKS = ['20260424', '20260518', '20260525', '20260529']

# 科室名称映射
DEPT_MAP = {
    '项目管理': '项目管理',
    '需求管理': '需求管理',
    '架构管理': '架构管理',
    '综合管理部': '综合管理部',
    '机构服务团队': '机构服务团队',
    '信息统计部': '信息统计部',
    '信息管理部': '信息管理部',
    '数据开发部': '数据开发部',
    '数据平台部': '数据平台部',
    '数据治理部': '数据治理部',
    '数据测试部': '数据测试部',
    '智能研发部-智能平台部': '智能平台部',
    '智能平台部': '智能平台部',
    '智能研发部-研发管理部': '研发管理部',
    '研发管理部': '研发管理部',
    '智能研发部-智能应用一部': '智能应用一部',
    '智能应用一部': '智能应用一部',
    '智能研发部-智能应用二部': '智能应用二部',
    '智能应用二部': '智能应用二部',
}

# 用户映射
USER_MAP = {
    '项目管理': {'id': 'xmgl', 'name': '项目管理'},
    '需求管理': {'id': 'xqgl', 'name': '需求管理'},
    '架构管理': {'id': 'jggl', 'name': '架构管理'},
    '综合管理部': {'id': '319915', 'name': '马胤'},
    '机构服务团队': {'id': '305069', 'name': '舒宝龙'},
    '信息统计部': {'id': '300523', 'name': '贺文军'},
    '信息管理部': {'id': '302390', 'name': '李焕彰'},
    '数据开发部': {'id': '319914', 'name': '顾恺'},
    '数据平台部': {'id': '306844', 'name': '吴证'},
    '数据治理部': {'id': '305249', 'name': '单曙兵'},
    '数据测试部': {'id': '306253', 'name': '杨萍'},
    '智能平台部': {'id': '307298', 'name': '胡申民'},
    '研发管理部': {'id': '304105', 'name': '刘异'},
    '智能应用一部': {'id': '302577', 'name': '杨晓彦'},
    '智能应用二部': {'id': '305393', 'name': '陈嘉琳'},
}


def detect_columns(headers):
    """智能匹配列名"""
    dept_col = plan_col = current_col = next_col = thoughts_col = ai_col = other_col = None

    for i, h in enumerate(headers):
        if h is None:
            continue
        h = str(h).strip()

        if dept_col is None and '科室' in h and '委员会' in h:
            dept_col = i
            continue

        if '工作内容' in h:
            current_col = i
            continue

        if '下周' in h:
            next_col = i
            continue

        if '管理心得' in h or (('心得' in h or '体会' in h) and 'AI' not in h and '推广' not in h):
            if thoughts_col is None:
                thoughts_col = i
            continue

        if 'AI' in h or '推广' in h:
            ai_col = i
            continue

        if '问题' in h or '风险' in h:
            other_col = i
            continue

        if '计划' in h and '下周' not in h and '工作' not in h:
            plan_col = i
            continue
        elif '计划' in h and '本周' in h and '下周' not in h and '工作' not in h:
            if plan_col is None:
                plan_col = i
                continue

    if plan_col is None:
        for i, h in enumerate(headers):
            if h and '计划' in str(h) and '下周' not in str(h):
                if current_col is None or i != current_col:
                    plan_col = i
                    break

    return dept_col, plan_col, current_col, next_col, thoughts_col, ai_col, other_col


def parse_plan_to_tasks(plan_text):
    """将计划文本解析为任务列表（兼容旧格式，前端会动态剥离前缀）"""
    if not plan_text:
        return []

    tasks = []
    lines = str(plan_text).strip().split('\n')

    for line in lines:
        line = line.strip()
        if not line:
            continue

        main_match = re.match(r'^(\d+[、.．）])\s*(.+)$', line)
        if main_match:
            tasks.append({'id': f'task-{len(tasks)}', 'text': line, 'checked': False})
        elif re.match(r'^[（(]\d+[）)]', line) or re.match(r'^\d+[）)]', line):
            tasks.append({'id': f'task-{len(tasks)}', 'text': line, 'checked': False})
        else:
            tasks.append({'id': f'task-{len(tasks)}', 'text': line, 'checked': False})

    return tasks


def parse_excel_files():
    """解析目标 Excel 文件，返回 report 列表"""
    all_reports = []

    for week_label in TARGET_WEEKS:
        filename = f'数据部工作周报-{week_label}.xlsx'
        filepath = os.path.join(EXCEL_DIR, filename)

        if not os.path.exists(filepath):
            print(f"⚠️  文件不存在: {filepath}")
            continue

        wb = openpyxl.load_workbook(filepath)
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        cols = detect_columns(headers)

        dept_col, plan_col, current_col, next_col, thoughts_col, ai_col, other_col = cols

        print(f"\n📄 解析 {filename} (sheet: {sheet_name})")
        print(f"   列: dept={dept_col}, plan={plan_col}, current={current_col}, next={next_col}, "
              f"thoughts={thoughts_col}, ai={ai_col}, other={other_col}")

        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            raw_dept = str(row[dept_col]).strip() if dept_col is not None and row[dept_col] else ''
            if not raw_dept or raw_dept in ['nan', 'None', '']:
                continue

            dept = DEPT_MAP.get(raw_dept, raw_dept)
            user = USER_MAP.get(dept, {'id': 'admin', 'name': '系统管理员'})

            plan_text = str(row[plan_col]).strip() if plan_col is not None and row[plan_col] else ''
            current_text = str(row[current_col]).strip() if current_col is not None and row[current_col] else ''
            next_text = str(row[next_col]).strip() if next_col is not None and row[next_col] else ''

            thoughts_parts = []
            if thoughts_col is not None and row[thoughts_col]:
                t = str(row[thoughts_col]).strip()
                if t and t not in ['无', '暂无', '']:
                    thoughts_parts.append(f"【本周管理心得】\n{t}")
            if ai_col is not None and row[ai_col]:
                t = str(row[ai_col]).strip()
                if t and t not in ['无', '暂无', '']:
                    thoughts_parts.append(f"【AI推广案例/心得】\n{t}")
            thoughts = '\n\n'.join(thoughts_parts)

            other = ''
            if other_col is not None and row[other_col]:
                o = str(row[other_col]).strip()
                if o and o not in ['无', '暂无', '']:
                    other = o

            report = {
                'id': f"{week_label}-{dept}-{idx}",
                'weekLabel': week_label,
                'dept': dept,
                'authorId': user['id'],
                'authorName': user['name'],
                'plan': plan_text,
                'content': parse_plan_to_tasks(current_text),
                'currentWork': current_text,
                'nextPlan': next_text,
                'thoughts': thoughts,
                'other': other,
            }
            all_reports.append(report)

    print(f"\n✅ 共解析出 {len(all_reports)} 条记录")
    return all_reports


def import_to_sqlite(reports, dry_run=True):
    """导入到 SQLite"""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    insert_count = 0
    update_count = 0
    now = datetime.now().isoformat()

    for report in reports:
        week_label = report['weekLabel']
        dept = report['dept']

        # 检查是否存在
        cursor.execute(
            "SELECT 1 FROM weekly_reports WHERE week_label = ? AND dept = ?",
            (week_label, dept)
        )
        exists = cursor.fetchone() is not None

        content_json = json.dumps(parse_plan_to_tasks(report["currentWork"]), ensure_ascii=False)

        if exists:
            if dry_run:
                print(f"[DRY-RUN UPDATE] {week_label} / {dept}")
            else:
                cursor.execute("""
                    UPDATE weekly_reports SET
                        plan = ?,
                        content = ?,
                        current_work = ?,
                        next_plan = ?,
                        thoughts = ?,
                        other = ?,
                        locked = 1,
                        updated_at = ?
                    WHERE week_label = ? AND dept = ?
                """, (
                    report['plan'],
                    content_json,
                    report['currentWork'],
                    report['nextPlan'],
                    report['thoughts'],
                    report['other'],
                    now,
                    week_label,
                    dept,
                ))
                update_count += 1
        else:
            if dry_run:
                print(f"[DRY-RUN INSERT] {week_label} / {dept}")
            else:
                cursor.execute("""
                    INSERT INTO weekly_reports
                    (id, week_label, dept, author_id, author_name, plan, content,
                     current_work, next_plan, thoughts, other, comments, submissions,
                     locked, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    report['id'],
                    week_label,
                    dept,
                    report['authorId'],
                    report['authorName'],
                    report['plan'],
                    content_json,
                    report['currentWork'],
                    report['nextPlan'],
                    report['thoughts'],
                    report['other'],
                    '[]',           # comments
                    '[]',           # submissions
                    1,              # locked = 1 (历史数据)
                    now,
                    now,
                ))
                insert_count += 1

    if not dry_run:
        conn.commit()
        print(f"\n✅ 导入完成: INSERT={insert_count}, UPDATE={update_count}")
    else:
        print(f"\n📋 预览完成: 将 INSERT={sum(1 for r in reports if not check_exists(conn, r))}, "
              f"UPDATE={sum(1 for r in reports if check_exists(conn, r))}")

    conn.close()


def check_exists(conn, report):
    cursor = conn.cursor()
    cursor.execute(
        "SELECT 1 FROM weekly_reports WHERE week_label = ? AND dept = ?",
        (report['weekLabel'], report['dept'])
    )
    return cursor.fetchone() is not None


if __name__ == '__main__':
    import sys
    dry = '--execute' not in sys.argv

    reports = parse_excel_files()

    if dry:
        print("\n" + "=" * 60)
        print("当前为预览模式 (dry-run)，不会写入数据库")
        print("确认无误后，执行: python3 scripts/import_new_weeks.py --execute")
        print("=" * 60)

    import_to_sqlite(reports, dry_run=dry)
